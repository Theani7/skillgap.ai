import os
from dotenv import load_dotenv
import google.generativeai as genai
import pypdf
import docx
from defusedxml import ElementTree as DET
import json
import logging

logger = logging.getLogger("resume-analyzer")
load_dotenv()

api_key = os.getenv("GEMINI_API_KEY", "").strip()
if api_key:
    logger.info("GEMINI_API_KEY is configured.")
else:
    logger.warning("GEMINI_API_KEY is not set. Gemini-powered features will be disabled, falling back to local hybrid parser.")
if not api_key:
    model = None
else:
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel('gemini-2.0-flash')

def extract_text_from_pdf(pdf_path: str) -> str:
    """Extracts raw text from a PDF file."""
    try:
        reader = pypdf.PdfReader(pdf_path)
        text = ""
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extracting PDF text: {e}")
        return ""

def extract_text_from_docx(docx_path: str) -> str:
    """Extracts raw text from a DOCX file."""
    try:
        doc = docx.Document(docx_path)
        text = ""
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extracting DOCX text: {e}")
        return ""

def parse_resume_with_gemini(file_path: str, target_role: str = None) -> dict:
    """
    Reads a PDF or DOCX, sends the text to Gemini, and extracts structured JSON containing
    the candidate's profile, education, experience, and skills.
    """
    if file_path.lower().endswith('.pdf'):
        raw_text = extract_text_from_pdf(file_path)
    elif file_path.lower().endswith('.docx'):
        raw_text = extract_text_from_docx(file_path)
    else:
        return {}
    
    if not raw_text.strip():
        return {}
        
    role_instruction = f"The candidate is applying for the target role of: '{target_role}'." if target_role else ""

    prompt = f"""
    You are an expert Technical Recruiter, AI Parser, and Career Coach.
    Carefully read the following resume text and extract the information into the EXACT JSON structure defined below.
    If a field is missing from the resume, return null or an empty array/string as appropriate.
    DO NOT return Markdown formatting. Return ONLY valid JSON.
    
    {role_instruction}
    If a target role is provided, you MUST evaluate the candidate's skills against standard industry requirements for that role.
    Determine what crucial skills the candidate is missing for that role, and return them as an array of strings in 'missing_skills'.
    Estimate how well the candidate's qualifications match the target role as an integer percentage (0-100) in 'match_score'. If no target role is explicitly provided, evaluate them based on their most prominent apparent field.

    CRITICAL — STRUCTURED EXPERIENCE & EDUCATION:
    You MUST populate 'experience_blocks' and 'education_blocks' as structured arrays.
    For each experience entry, extract: title, company, start_date, end_date, bullets (array of strings).
    For each education entry, extract: degree, institution, year.
    These structured fields are required by the frontend. Do NOT leave them empty if the resume contains experience or education.

    CRITICAL NEW INSTRUCTION - ROADMAP:
    Generate a highly personalized, 4-phase chronological learning path aimed at helping the user master their 'missing_skills' and secure the target role (or next logical career step). 
    Return this in the 'roadmap' array. Each phase must include specific 'action_items'. Make the advice practical and actionable.

    RESUME TEXT:
    {raw_text}
    """

    # We enforce a strict JSON schema so that we can reliably parse the output
    # in our main application, perfectly replacing the old pyresparser dictionary.
    response_schema = {
        "type": "OBJECT",
        "properties": {
            "name": {"type": "STRING"},
            "email": {"type": "STRING"},
            "mobile_number": {"type": "STRING"},
            "skills": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
                "description": "A comprehensive list of all technical and soft skills."
            },
            "education": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
                "description": "List of degrees, universities, and graduation years (flat strings for backward compat)."
            },
            "education_blocks": {
                "type": "ARRAY",
                "items": {
                    "type": "OBJECT",
                    "properties": {
                        "degree": {"type": "STRING", "description": "e.g. B.S., M.S., MBA, Ph.D."},
                        "institution": {"type": "STRING", "description": "University or college name"},
                        "year": {"type": "STRING", "description": "Graduation year, e.g. 2020"}
                    },
                    "required": ["degree", "institution", "year"]
                },
                "description": "Structured education entries."
            },
            "experience": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
                "description": "List of roles, companies, and descriptions (flat strings for backward compat)."
            },
            "experience_blocks": {
                "type": "ARRAY",
                "items": {
                    "type": "OBJECT",
                    "properties": {
                        "title": {"type": "STRING", "description": "Job title"},
                        "company": {"type": "STRING", "description": "Company name"},
                        "start_date": {"type": "STRING", "description": "e.g. Jan 2020"},
                        "end_date": {"type": "STRING", "description": "e.g. Present, Dec 2023"},
                        "bullets": {
                            "type": "ARRAY",
                            "items": {"type": "STRING"},
                            "description": "Bullet-point accomplishments for this role"
                        }
                    },
                    "required": ["title", "company", "start_date", "end_date", "bullets"]
                },
                "description": "Structured experience entries with title, company, dates, and bullet points."
            },
            "company_names": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
                "description": "List of companies the person has worked for."
            },
            "designation": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
                "description": "List of job titles the person has held."
            },
            "missing_skills": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
                "description": "Prioritized list of crucial skills the candidate lacks for the target role."
            },
            "match_score": {
                "type": "INTEGER",
                "description": "Percentage representing alignment with target role."
            },
            "no_of_pages": {"type": "INTEGER", "description": "Estimate 1 if unknown."},
            "roadmap": {
                "type": "ARRAY",
                "items": {
                    "type": "OBJECT",
                    "properties": {
                        "step": {"type": "INTEGER", "description": "Phase number (1-4)"},
                        "title": {"type": "STRING", "description": "Title of the phase"},
                        "duration": {"type": "STRING", "description": "Estimated duration (e.g., '3-4 weeks')"},
                        "skills": {
                            "type": "ARRAY",
                            "items": {"type": "STRING"},
                            "description": "Skills focused on in this phase"
                        },
                        "action_items": {
                            "type": "ARRAY",
                            "items": {"type": "STRING"},
                            "description": "Specific, actionable steps to master the skills in this phase"
                        }
                    },
                    "required": ["step", "title", "duration", "skills", "action_items"]
                },
                "description": "A 4-phase personalized chronological learning path."
            }
        },
        "required": ["name", "email", "skills", "education", "experience", "experience_blocks", "education_blocks", "roadmap"]
    }

    try:
        response = model.generate_content(
            prompt,
            generation_config=genai.types.GenerationConfig(
                response_mime_type="application/json",
                response_schema=response_schema,
                temperature=0.1,
            ),
        )
        
        if not response or not response.text:
            logger.error("Gemini returned empty response")
            return {}
        
        parsed_data = json.loads(response.text)
        
        if not isinstance(parsed_data, dict):
            logger.error("Gemini response is not a valid JSON object")
            return {}
        
        try:
            if file_path.lower().endswith('.pdf'):
                reader = pypdf.PdfReader(file_path)
                parsed_data["no_of_pages"] = len(reader.pages)
            else:
                parsed_data["no_of_pages"] = 1
        except Exception as page_error:
            logger.warning(f"Could not determine page count: {page_error}")
            parsed_data["no_of_pages"] = 1
            
        return parsed_data

    except json.JSONDecodeError as json_err:
        logger.error(f"Failed to parse Gemini JSON response: {json_err}")
        return {}
    except Exception as e:
        logger.error(f"Error during Gemini extraction: {e}")
        return {}


def rewrite_resume_with_gemini(resume_data: dict, target_role: str = None) -> dict:
    role = target_role or "Target Role"
    prompt = f"""
    You are an expert resume editor.
    Rewrite resume experience bullets for the role: {role}.
    Return JSON only with this schema:
    {{
      "target_role": "string",
      "rewritten_bullets": [
        {{"before": "string", "after": "string"}}
      ]
    }}
    Use action verbs and quantified outcomes where possible.
    Input resume JSON:
    {json.dumps(resume_data)}
    """
    response_schema = {
        "type": "OBJECT",
        "properties": {
            "target_role": {"type": "STRING"},
            "rewritten_bullets": {
                "type": "ARRAY",
                "items": {
                    "type": "OBJECT",
                    "properties": {
                        "before": {"type": "STRING"},
                        "after": {"type": "STRING"}
                    },
                    "required": ["before", "after"]
                }
            }
        },
        "required": ["target_role", "rewritten_bullets"]
    }
    try:
        response = model.generate_content(
            contents=prompt,
            generation_config=genai.types.GenerationConfig(
                response_mime_type="application/json",
                response_schema=response_schema,
                temperature=0.2,
            ),
        )
        if not response or not response.text:
            logger.error("Gemini returned empty response for resume rewrite")
            return {}
        return json.loads(response.text)
    except json.JSONDecodeError as json_err:
        logger.error(f"Failed to parse Gemini resume rewrite response: {json_err}")
        return {}
    except Exception as e:
        logger.error(f"Error during Gemini resume rewrite: {e}")
        return {}


def generate_cover_letter_with_gemini(profile: dict, job_description: str, company: str, role: str) -> str:
    name = profile.get('name', 'Candidate')
    email = profile.get('email', '')
    background = profile.get('background', '')
    skills = profile.get('skills', '')
    achievement = profile.get('achievement', '')
    why_company = profile.get('why_company', '')
    
    if not role or len(str(role).strip()) < 4:
        return "ERROR: Please provide a valid job title before generating."
    
    prompt = f"""You are an expert career coach and professional cover letter writer.

Write a polished, personalized, and compelling cover letter using ONLY the information provided below. 

STRICT RULES:
- Never use filler phrases like "problem solving", "team player", or "passionate about"
- Never leave placeholders, blanks, or template text in the output
- Every sentence must add specific value — no fluff
- Do not repeat the same point twice
- Use a confident, professional, and human tone
- Output ONLY the cover letter text, no explanations or commentary

STRUCTURE (3 paragraphs):
1. Opening — State the role and company, and make a strong first impression with a specific reason you're a great fit
2. Middle — Highlight 2-3 relevant skills tied to the role, and include the achievement with impact
3. Closing — Express enthusiasm for this specific company, and invite next steps confidently

INPUT DATA:
- Applicant Name: {name}
- Email: {email}
- Target Role: {role}
- Company: {company}
- Years of Experience / Background: {background}
- Key Skills: {skills}
- A Key Achievement (with metric if possible): {achievement}
- Why this company: {why_company}
- Job Description: {job_description}

If the Target Role looks invalid or too short (under 4 characters), write:
"ERROR: Please provide a valid job title before generating."

If Job Description is provided, tailor the letter closely to the listed requirements.
If Job Description is empty, write a strong general letter based on the role and company.

Now write the cover letter:"""
    
    try:
        response = model.generate_content(
            contents=prompt,
            generation_config=genai.types.GenerationConfig(
                temperature=0.2,
            ),
        )
        if not response or not response.text:
            logger.error("Gemini returned empty response for cover letter generation")
            return ""
        return (response.text or "").strip()
    except Exception as e:
        logger.error(f"Error during Gemini cover letter generation: {e}")
        return ""

def simulate_interview_turn(resume_data: dict, target_role: str, chat_history: list) -> dict:
    """
    Simulates a turn in an AI-driven interview.
    Generates a follow-up question based on resume data and history.
    """
    if model is None:
        logger.warning("simulate_interview_turn called but Gemini model is not available")
        return {
            "feedback": "Thanks for your answer.",
            "question": "Can you elaborate on your experience with key projects?",
            "is_finished": False
        }

    history_str = "\n".join([f"{m['role']}: {m['content']}" for m in chat_history])
    
    prompt = f"""
    You are an expert Technical Interviewer for the role of {target_role}.
    You are interviewing a candidate with the following resume details:
    {json.dumps(resume_data)}
    
    Current Interview History:
    {history_str}
    
    RULES:
    1. If this is the start (history is empty), introduce yourself briefly and ask the first question.
    2. Ask deep, behavioral or technical questions based on the candidate's actual projects or skills.
    3. Provide brief, encouraging feedback on their last answer before asking the next question.
    4. Return ONLY a JSON object with this schema:
    {{
      "feedback": "Your coaching/feedback on the last answer",
      "question": "Your next interview question",
      "is_finished": boolean
    }}
    5. After 3-4 turns, wrap up the interview and set is_finished to true.
    """
    
    try:
        response = model.generate_content(
            contents=prompt,
            generation_config=genai.types.GenerationConfig(
                response_mime_type="application/json",
                temperature=0.7,
            ),
        )
        return json.loads(response.text)
    except Exception as e:
        logger.error(f"Interview simulation failed: {e}")
        return {
            "feedback": "Interesting perspective.",
            "question": "Can you tell me more about your most challenging technical project?",
            "is_finished": False
        }

